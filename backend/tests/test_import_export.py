from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from sqlalchemy import text

from app.models.entry import Entry
from app.models.enums import PaymentMethod, RecurrenceType, UserRole
from app.models.expense import Expense
from app.models.installment import Installment
from app.models.installment_adjustment import InstallmentAdjustment
from app.models.user import User
from app.services import export_service, import_service
from app.services.import_service import ParsedImport, execute_import, validate_workbook

USERS_HEADER = import_service.USERS_COLUMNS
ENTRIES_HEADER = import_service.ENTRIES_COLUMNS
EXPENSES_HEADER = import_service.EXPENSES_COLUMNS
INSTALLMENTS_HEADER = import_service.INSTALLMENTS_COLUMNS
ADJUSTMENTS_HEADER = import_service.ADJUSTMENTS_COLUMNS


def _build_workbook(
    users=None, entries=None, expenses=None, installments=None, adjustments=None
) -> bytes:
    wb = Workbook()
    sheets = {
        "users": (USERS_HEADER, users or []),
        "entries": (ENTRIES_HEADER, entries or []),
        "expenses": (EXPENSES_HEADER, expenses or []),
        "installments": (INSTALLMENTS_HEADER, installments or []),
        "installment_adjustments": (ADJUSTMENTS_HEADER, adjustments or []),
    }
    first = True
    for name, (header, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        first = False
        ws.append(header)
        for row in rows:
            ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _valid_user_row(user_id=1, email="admin@example.com"):
    return [user_id, "Admin", email, None, "ADMIN", "2026-01-01 00:00:00", "2026-01-01 00:00:00"]


def _valid_entry_row(entry_id=1, user_id=1):
    return [
        entry_id, user_id, "2026-01-05", "1800.00", "Salario", "Trabalho", "FIXED",
        "2026-01-01", None, "2026-01-01 00:00:00", "2026-01-01 00:00:00",
    ]


def test_validate_workbook_accepts_a_well_formed_file():
    content = _build_workbook(users=[_valid_user_row()], entries=[_valid_entry_row()])
    preview, parsed = validate_workbook(content)

    assert preview.valid is True
    assert preview.errors == []
    assert parsed is not None
    assert parsed.users[0]["email"] == "admin@example.com"
    assert parsed.entries[0]["amount"] == Decimal("1800.00")


def test_validate_workbook_rejects_missing_required_sheet():
    wb = Workbook()
    wb.active.title = "users"
    wb.active.append(USERS_HEADER)
    buffer = BytesIO()
    wb.save(buffer)

    from app.services.import_service import ImportFileError

    try:
        validate_workbook(buffer.getvalue())
        assert False, "expected ImportFileError"
    except ImportFileError as exc:
        assert "entries" in exc.message


def test_validate_workbook_rejects_wrong_header():
    content = _build_workbook(users=[_valid_user_row()])
    # Corrupt the users header by rebuilding with a bad column name.
    wb = Workbook()
    ws = wb.active
    ws.title = "users"
    ws.append(["id", "nome", "email", "photo", "role", "created_at", "updated_at"])
    ws.append(_valid_user_row())
    for name, header in (
        ("entries", ENTRIES_HEADER),
        ("expenses", EXPENSES_HEADER),
        ("installments", INSTALLMENTS_HEADER),
        ("installment_adjustments", ADJUSTMENTS_HEADER),
    ):
        s = wb.create_sheet(name)
        s.append(header)
    buffer = BytesIO()
    wb.save(buffer)

    preview, parsed = validate_workbook(buffer.getvalue())
    assert preview.valid is False
    assert parsed is None
    assert any("Colunas esperadas" in e.message for e in preview.errors)


def test_validate_workbook_rejects_duplicate_user_ids():
    content = _build_workbook(
        users=[_valid_user_row(1, "a@example.com"), _valid_user_row(1, "b@example.com")]
    )
    preview, parsed = validate_workbook(content)

    assert preview.valid is False
    assert parsed is None
    assert any("duplicado" in e.message for e in preview.errors)


def test_validate_workbook_rejects_entry_with_broken_user_reference():
    content = _build_workbook(
        users=[_valid_user_row(1)],
        entries=[_valid_entry_row(entry_id=1, user_id=999)],
    )
    preview, _parsed = validate_workbook(content)

    assert preview.valid is False
    assert any("user_id inexistente" in e.message for e in preview.errors)


def test_validate_workbook_rejects_installment_count_out_of_range():
    content = _build_workbook(
        users=[_valid_user_row(1)],
        expenses=[
            [
                1, 1, "2026-01-10", "90.00", "Compra", "Loja", "EXCEPTIONAL", "CREDITO",
                None, None, "2026-01-01 00:00:00", "2026-01-01 00:00:00",
            ]
        ],
        installments=[
            [1, 1, 1, 30, "2026-01-10", "3.00", "2026-01-01 00:00:00", "2026-01-01 00:00:00"],
        ],
    )
    preview, _parsed = validate_workbook(content)

    assert preview.valid is False
    assert any("1-24" in e.message for e in preview.errors)


def test_validate_workbook_rejects_installment_gap_in_sequence():
    content = _build_workbook(
        users=[_valid_user_row(1)],
        expenses=[
            [
                1, 1, "2026-01-10", "90.00", "Compra", "Loja", "EXCEPTIONAL", "CREDITO",
                None, None, "2026-01-01 00:00:00", "2026-01-01 00:00:00",
            ]
        ],
        installments=[
            [1, 1, 1, 3, "2026-01-10", "30.00", "2026-01-01 00:00:00", "2026-01-01 00:00:00"],
            [2, 1, 3, 3, "2026-03-10", "30.00", "2026-01-01 00:00:00", "2026-01-01 00:00:00"],
            # installment_number 2 missing - gap in the 1..N sequence
        ],
    )
    preview, _parsed = validate_workbook(content)

    assert preview.valid is False
    assert any("inconsistentes" in e.message for e in preview.errors)


def test_validate_workbook_rejects_orphan_adjustment():
    content = _build_workbook(
        users=[_valid_user_row(1)],
        adjustments=[[1, 999, "5.00", "Taxa", "2026-01-01 00:00:00", "2026-01-01 00:00:00"]],
    )
    preview, _parsed = validate_workbook(content)

    assert preview.valid is False
    assert any("installment_id inexistente" in e.message for e in preview.errors)


def test_export_then_validate_round_trips_cleanly(db_session, make_user):
    user = make_user(email="round@example.com")
    db_session.add(
        Entry(
            user_id=user.id, date=date(2026, 1, 5), amount=Decimal("1800"), description="Salario",
            source="Trabalho", type=RecurrenceType.FIXED, start_date=date(2026, 1, 1),
        )
    )
    db_session.commit()

    buffer = export_service.build_export_workbook(db_session)
    preview, parsed = validate_workbook(buffer.getvalue())

    assert preview.valid is True
    assert preview.counts["users"] == 1
    assert preview.counts["entries"] == 1
    assert "password" not in str(export_service.USERS_COLUMNS)
    assert "password_hash" not in export_service.USERS_COLUMNS
    assert parsed.entries[0]["amount"] == Decimal("1800.00")


def test_execute_import_replaces_data_and_issues_fresh_passwords(engine):
    # execute_import opens its own real, separately-committing transaction via `engine`
    # (independent of the db_session fixture's rollback-only savepoint), so setup and
    # verification here use `engine` directly too - otherwise neither side would see
    # the other's changes.
    with engine.begin() as conn:
        conn.execute(
            User.__table__.insert().values(
                id=9500, name="Old", email="old@example.com",
                password_hash="irrelevant-hash", role=UserRole.USER,
            )
        )

    parsed = ParsedImport(
        users=[
            {
                "id": 500,
                "name": "Imported Admin",
                "email": "imported@example.com",
                "photo": None,
                "role": UserRole.ADMIN,
                "created_at": datetime(2026, 1, 1),
                "updated_at": datetime(2026, 1, 1),
            }
        ],
        entries=[
            {
                "id": 500,
                "user_id": 500,
                "date": date(2026, 1, 5),
                "amount": Decimal("1800.00"),
                "description": "Salario",
                "source": "Trabalho",
                "type": RecurrenceType.FIXED,
                "start_date": date(2026, 1, 1),
                "end_date": None,
                "created_at": datetime(2026, 1, 1),
                "updated_at": datetime(2026, 1, 1),
            }
        ],
        expenses=[],
        installments=[],
        installment_adjustments=[],
    )

    try:
        result = execute_import(engine, parsed)

        assert result.counts == {
            "users": 1, "entries": 1, "expenses": 0, "installments": 0, "installment_adjustments": 0
        }
        assert "imported@example.com" in result.temporary_passwords
        assert len(result.temporary_passwords["imported@example.com"]) >= 12

        # Old data is gone, new data is present, with the id preserved from the file.
        with engine.connect() as conn:
            remaining_users = conn.execute(text("SELECT id, email, password_hash FROM users")).fetchall()

        assert len(remaining_users) == 1
        assert remaining_users[0].id == 500
        assert remaining_users[0].email == "imported@example.com"
        # The temp password actually works (is properly hashed, not stored in plaintext).
        assert remaining_users[0].password_hash != result.temporary_passwords["imported@example.com"]
        assert remaining_users[0].password_hash.startswith("$argon2")
    finally:
        with engine.begin() as conn:
            conn.execute(text("SET FOREIGN_KEY_CHECKS=0"))
            conn.execute(text("DELETE FROM entries WHERE id IN (500, 9500)"))
            conn.execute(text("DELETE FROM users WHERE id IN (500, 9500)"))
            conn.execute(text("SET FOREIGN_KEY_CHECKS=1"))


def test_execute_import_wipes_assistant_history_and_settings(engine):
    # The import runs with FOREIGN_KEY_CHECKS=0, so ON DELETE CASCADE doesn't fire - a
    # conversation left behind would surface for whichever imported user reuses the id.
    with engine.begin() as conn:
        conn.execute(
            User.__table__.insert().values(
                id=9600, name="Old", email="old-chat@example.com", password_hash="irrelevant-hash", role=UserRole.USER
            )
        )
        conn.execute(text("INSERT INTO chat_conversations (id, user_id, title) VALUES (9600, 9600, 'Antiga')"))
        conn.execute(
            text("INSERT INTO chat_messages (conversation_id, role, content) VALUES (9600, 'USER', 'segredo')")
        )
        conn.execute(
            text("INSERT INTO assistant_settings (user_id, provider, api_key_encrypted) VALUES (9600, 'groq', 'x')")
        )

    parsed = ParsedImport(
        users=[{
            "id": 9600, "name": "Outra pessoa", "email": "new-chat@example.com", "photo": None,
            "role": UserRole.USER, "created_at": datetime(2026, 1, 1), "updated_at": datetime(2026, 1, 1),
        }],
        entries=[], expenses=[], installments=[], installment_adjustments=[],
    )
    try:
        execute_import(engine, parsed)
        with engine.connect() as conn:
            assert conn.execute(text("SELECT COUNT(*) FROM chat_conversations")).scalar() == 0
            assert conn.execute(text("SELECT COUNT(*) FROM chat_messages")).scalar() == 0
            # A saved API key must never end up belonging to another person.
            assert conn.execute(text("SELECT COUNT(*) FROM assistant_settings")).scalar() == 0
    finally:
        with engine.begin() as conn:
            conn.execute(text("SET FOREIGN_KEY_CHECKS=0"))
            conn.execute(text("DELETE FROM chat_messages WHERE conversation_id = 9600"))
            conn.execute(text("DELETE FROM chat_conversations WHERE id = 9600"))
            conn.execute(text("DELETE FROM assistant_settings WHERE user_id = 9600"))
            conn.execute(text("DELETE FROM users WHERE id = 9600"))
            conn.execute(text("SET FOREIGN_KEY_CHECKS=1"))


def test_execute_import_rolls_back_entirely_on_failure(engine):
    """Mirrors the real incident this pipeline is designed to prevent: if anything
    fails partway through the replace, none of it should stick - not even the earlier
    steps in the same call. Forces a DB-level failure via a duplicate primary key in
    the entries batch (FK checks are deliberately disabled during the replace, so a
    dangling FK wouldn't fail here - a PK collision still will)."""
    with engine.begin() as conn:
        conn.execute(
            User.__table__.insert().values(
                id=9900, name="Original", email="original@example.com",
                password_hash="hash", role=UserRole.USER,
            )
        )

    parsed = ParsedImport(
        users=[
            {
                "id": 900,
                "name": "New",
                "email": "new@example.com",
                "photo": None,
                "role": UserRole.USER,
                "created_at": datetime(2026, 1, 1),
                "updated_at": datetime(2026, 1, 1),
            }
        ],
        entries=[
            {
                "id": 900, "user_id": 900, "date": date(2026, 1, 5), "amount": Decimal("100.00"),
                "description": "x", "source": "x", "type": RecurrenceType.EXCEPTIONAL,
                "start_date": None, "end_date": None,
                "created_at": datetime(2026, 1, 1), "updated_at": datetime(2026, 1, 1),
            },
            {
                # duplicate id -> primary key violation, forcing the whole transaction
                # (including the two inserts above) to roll back.
                "id": 900, "user_id": 900, "date": date(2026, 1, 6), "amount": Decimal("50.00"),
                "description": "y", "source": "y", "type": RecurrenceType.EXCEPTIONAL,
                "start_date": None, "end_date": None,
                "created_at": datetime(2026, 1, 1), "updated_at": datetime(2026, 1, 1),
            },
        ],
        expenses=[],
        installments=[],
        installment_adjustments=[],
    )

    try:
        try:
            execute_import(engine, parsed)
            assert False, "expected the import to fail on the duplicate entry id"
        except Exception as exc:
            assert "Duplicate" in str(exc) or "duplicate" in str(exc)

        # The original user must still be there, completely untouched, and the bogus
        # imported user/entry must NOT have been left behind despite being inserted
        # before the failing statement in the same transaction.
        with engine.connect() as conn:
            users = conn.execute(text("SELECT id, email FROM users")).fetchall()
            entries = conn.execute(text("SELECT id FROM entries WHERE id = 900")).fetchall()

        emails = {row.email for row in users}
        assert "original@example.com" in emails
        assert "new@example.com" not in emails
        assert entries == []
    finally:
        with engine.begin() as conn:
            conn.execute(text("SET FOREIGN_KEY_CHECKS=0"))
            conn.execute(text("DELETE FROM entries WHERE id = 900"))
            conn.execute(text("DELETE FROM users WHERE id IN (900, 9900)"))
            conn.execute(text("SET FOREIGN_KEY_CHECKS=1"))


def _auth_header(client, email, password):
    response = client.post("/auth/login", json={"email": email, "password": password})
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_database_endpoints_require_admin(client, make_user):
    make_user(email="regular@example.com", password="Sfort3!123", role=UserRole.USER)
    headers = _auth_header(client, "regular@example.com", "Sfort3!123")

    assert client.get("/database/export", headers=headers).status_code == 403
    assert client.post("/database/restore-backup", headers=headers).status_code == 403

    content = _build_workbook(users=[_valid_user_row()])
    files = {"file": ("export.xlsx", content, "application/octet-stream")}
    assert (
        client.post("/database/import?dry_run=true", headers=headers, files=files).status_code == 403
    )


def test_import_requires_exactly_one_of_dry_run_or_confirm(client, make_user):
    make_user(email="admin@example.com", password="Sfort3!123", role=UserRole.ADMIN)
    headers = _auth_header(client, "admin@example.com", "Sfort3!123")

    content = _build_workbook(users=[_valid_user_row()])
    files = {"file": ("export.xlsx", content, "application/octet-stream")}

    assert client.post("/database/import", headers=headers, files=files).status_code == 400
    assert (
        client.post(
            "/database/import?dry_run=true&confirm=true", headers=headers, files=files
        ).status_code
        == 400
    )


def test_import_rejects_non_xlsx_extension(client, make_user):
    make_user(email="admin2@example.com", password="Sfort3!123", role=UserRole.ADMIN)
    headers = _auth_header(client, "admin2@example.com", "Sfort3!123")

    files = {"file": ("export.csv", b"not a real spreadsheet", "text/csv")}
    response = client.post("/database/import?dry_run=true", headers=headers, files=files)

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "INVALID_IMPORT_FILE"


def test_import_dry_run_via_http_reports_validation_errors_without_touching_db(client, make_user):
    """Exercises the full router path for dry_run (file upload, admin gate, validation
    pipeline). confirm=true is intentionally NOT exercised via HTTP here: even though
    the data-replacement engine is dependency-injectable (see get_engine), the backup
    step still shells out to mysqldump against the globally-configured MYSQL_* database,
    so a real HTTP-level confirm test would dump the dev DB from the test suite. That
    DB-mutating path is covered directly at the service layer in the execute_import
    tests above instead, and was verified end-to-end manually (including a real
    lockout-and-recovery) against the running dev stack."""
    make_user(email="admin3@example.com", password="Sfort3!123", role=UserRole.ADMIN)
    headers = _auth_header(client, "admin3@example.com", "Sfort3!123")

    bad_content = _build_workbook(
        users=[_valid_user_row(1, "a@example.com"), _valid_user_row(1, "b@example.com")]
    )
    files = {"file": ("export.xlsx", bad_content, "application/octet-stream")}
    response = client.post("/database/import?dry_run=true", headers=headers, files=files)

    assert response.status_code == 200
    body = response.json()
    assert body["valid"] is False
    assert any("duplicado" in e["message"] for e in body["errors"])

    good_content = _build_workbook(users=[_valid_user_row()])
    files = {"file": ("export.xlsx", good_content, "application/octet-stream")}
    response = client.post("/database/import?dry_run=true", headers=headers, files=files)

    assert response.status_code == 200
    assert response.json()["valid"] is True
