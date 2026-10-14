"""Validates and executes a full-database xlsx import (spec sections 25 & 27).

Two-step flow, enforced by the router:
  1. `validate_workbook()` - parses + runs every check below, never touches the DB.
  2. `execute_import()` - only called after an admin explicitly confirms a valid
     preview; takes a mysqldump backup first (see backup_service), then replaces all
     data inside a single transaction. Any failure rolls back and leaves the DB
     untouched; the backup remains available for manual restore either way.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.auth.security import generate_temp_password, hash_password
from app.models.entry import Entry
from app.models.enums import PaymentMethod, RecurrenceType, UserRole
from app.models.expense import Expense
from app.models.installment import Installment
from app.models.installment_adjustment import InstallmentAdjustment
from app.models.user import User
from app.utils.errors import AppError

REQUIRED_SHEETS = ["users", "entries", "expenses", "installments", "installment_adjustments"]

USERS_COLUMNS = ["id", "name", "email", "photo", "role", "created_at", "updated_at"]
ENTRIES_COLUMNS = [
    "id", "user_id", "date", "amount", "description", "source", "type",
    "start_date", "end_date", "created_at", "updated_at",
]
EXPENSES_COLUMNS = [
    "id", "user_id", "date", "amount", "description", "source", "type", "payment_method",
    "start_date", "end_date", "created_at", "updated_at",
]
INSTALLMENTS_COLUMNS = [
    "id", "expense_id", "installment_number", "total_installments", "due_date",
    "base_amount", "advanced_to_year", "advanced_to_month", "created_at", "updated_at",
]
ADJUSTMENTS_COLUMNS = ["id", "installment_id", "amount", "description", "created_at", "updated_at"]

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ImportFileError(AppError):
    status_code = 400
    code = "INVALID_IMPORT_FILE"


@dataclass
class ValidationIssue:
    sheet: str
    row: int | None
    message: str

    def to_dict(self) -> dict:
        return {"sheet": self.sheet, "row": self.row, "message": self.message}


@dataclass
class ParsedImport:
    users: list[dict] = field(default_factory=list)
    entries: list[dict] = field(default_factory=list)
    expenses: list[dict] = field(default_factory=list)
    installments: list[dict] = field(default_factory=list)
    installment_adjustments: list[dict] = field(default_factory=list)


@dataclass
class ImportPreview:
    valid: bool
    errors: list[ValidationIssue]
    warnings: list[ValidationIssue]
    counts: dict[str, int]


# ---- typed cell parsing ----------------------------------------------------------


def _parse_int(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().lstrip("-").isdigit():
        return int(value)
    return None


def _parse_decimal(value) -> Decimal | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    if isinstance(value, str):
        try:
            return Decimal(value.strip()).quantize(Decimal("0.01"))
        except InvalidOperation:
            return None
    return None


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _parse_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _parse_str(value) -> str | None:
    if value is None:
        return None
    text_value = str(value).strip()
    return text_value if text_value else None


def _parse_enum(value, enum_cls):
    if value is None:
        return None
    text_value = str(value).strip().upper()
    try:
        return enum_cls(text_value)
    except ValueError:
        return None


# ---- workbook parsing --------------------------------------------------------------


def parse_workbook(file_bytes: bytes) -> dict[str, tuple[list, list[list]]]:
    """Returns {sheet_name: (header_row, data_rows)}."""
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ImportFileError(f"Não foi possível ler o arquivo como planilha .xlsx: {exc}") from exc

    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        raise ImportFileError(f"Abas obrigatórias ausentes: {', '.join(missing)}.")

    result = {}
    for name in REQUIRED_SHEETS:
        rows = list(wb[name].iter_rows(values_only=True))
        header, data = (rows[0], rows[1:]) if rows else ([], [])
        result[name] = (list(header), data)
    return result


def _check_header(ws_name: str, header: list, expected: list[str], issues: list[ValidationIssue]) -> bool:
    actual = [str(h).strip() if h is not None else "" for h in header]
    if actual != expected:
        issues.append(
            ValidationIssue(
                ws_name, None, f"Colunas esperadas {expected}, encontradas {actual}."
            )
        )
        return False
    return True


# ---- per-sheet validation -----------------------------------------------------------


def _validate_users(rows: list[list], issues: list[ValidationIssue]) -> tuple[list[dict], set[int]]:
    parsed: list[dict] = []
    seen_ids: set[int] = set()
    seen_emails: set[str] = set()

    for i, row in enumerate(rows, start=2):
        values = dict(zip(USERS_COLUMNS, row))
        row_id = _parse_int(values.get("id"))
        if row_id is None:
            issues.append(ValidationIssue("users", i, "id inválido ou ausente."))
            continue
        if row_id in seen_ids:
            issues.append(ValidationIssue("users", i, f"id duplicado: {row_id}."))
            continue
        seen_ids.add(row_id)

        name = _parse_str(values.get("name"))
        if name is None:
            issues.append(ValidationIssue("users", i, "name é obrigatório."))
            continue

        email = _parse_str(values.get("email"))
        if email is None or not EMAIL_RE.match(email):
            issues.append(ValidationIssue("users", i, f"email inválido: {values.get('email')!r}."))
            continue
        if email.lower() in seen_emails:
            issues.append(ValidationIssue("users", i, f"email duplicado: {email}."))
            continue
        seen_emails.add(email.lower())

        role = _parse_enum(values.get("role"), UserRole)
        if role is None:
            issues.append(ValidationIssue("users", i, f"role inválida: {values.get('role')!r}."))
            continue

        created_at = _parse_datetime(values.get("created_at")) or datetime.utcnow()
        updated_at = _parse_datetime(values.get("updated_at")) or datetime.utcnow()

        parsed.append(
            {
                "id": row_id,
                "name": name,
                "email": email,
                "photo": _parse_str(values.get("photo")),
                "role": role,
                "created_at": created_at,
                "updated_at": updated_at,
            }
        )

    return parsed, seen_ids


def _validate_recurrence_row(
    sheet: str, i: int, values: dict, columns: list[str], issues: list[ValidationIssue],
    valid_user_ids: set[int], seen_ids: set[int],
) -> dict | None:
    row_id = _parse_int(values.get("id"))
    if row_id is None:
        issues.append(ValidationIssue(sheet, i, "id inválido ou ausente."))
        return None
    if row_id in seen_ids:
        issues.append(ValidationIssue(sheet, i, f"id duplicado: {row_id}."))
        return None

    user_id = _parse_int(values.get("user_id"))
    if user_id is None or user_id not in valid_user_ids:
        issues.append(ValidationIssue(sheet, i, f"user_id inexistente: {values.get('user_id')!r}."))
        return None

    row_date = _parse_date(values.get("date"))
    if row_date is None:
        issues.append(ValidationIssue(sheet, i, "date inválida."))
        return None

    amount = _parse_decimal(values.get("amount"))
    if amount is None or amount <= 0:
        issues.append(ValidationIssue(sheet, i, f"amount inválido ou não positivo: {values.get('amount')!r}."))
        return None

    description = _parse_str(values.get("description"))
    source = _parse_str(values.get("source"))
    if description is None or source is None:
        issues.append(ValidationIssue(sheet, i, "description e source são obrigatórios."))
        return None

    row_type = _parse_enum(values.get("type"), RecurrenceType)
    if row_type is None:
        issues.append(ValidationIssue(sheet, i, f"type inválido: {values.get('type')!r}."))
        return None

    start_date = _parse_date(values.get("start_date")) if values.get("start_date") is not None else None
    end_date = _parse_date(values.get("end_date")) if values.get("end_date") is not None else None

    if row_type == RecurrenceType.FIXED and start_date is None:
        issues.append(ValidationIssue(sheet, i, "start_date é obrigatório para type=FIXED."))
        return None
    if end_date is not None and start_date is not None and end_date < start_date:
        issues.append(ValidationIssue(sheet, i, "end_date anterior a start_date."))
        return None

    seen_ids.add(row_id)
    created_at = _parse_datetime(values.get("created_at")) or datetime.utcnow()
    updated_at = _parse_datetime(values.get("updated_at")) or datetime.utcnow()

    return {
        "id": row_id,
        "user_id": user_id,
        "date": row_date,
        "amount": amount,
        "description": description,
        "source": source,
        "type": row_type,
        "start_date": start_date,
        "end_date": end_date,
        "created_at": created_at,
        "updated_at": updated_at,
    }


def _validate_entries(rows: list[list], issues: list[ValidationIssue], valid_user_ids: set[int]):
    parsed, seen_ids = [], set()
    for i, row in enumerate(rows, start=2):
        values = dict(zip(ENTRIES_COLUMNS, row))
        result = _validate_recurrence_row("entries", i, values, ENTRIES_COLUMNS, issues, valid_user_ids, seen_ids)
        if result is not None:
            parsed.append(result)
    return parsed, seen_ids


def _validate_expenses(rows: list[list], issues: list[ValidationIssue], valid_user_ids: set[int]):
    parsed, seen_ids = [], set()
    for i, row in enumerate(rows, start=2):
        values = dict(zip(EXPENSES_COLUMNS, row))
        result = _validate_recurrence_row(
            "expenses", i, values, EXPENSES_COLUMNS, issues, valid_user_ids, seen_ids
        )
        if result is None:
            continue

        payment_method = _parse_enum(values.get("payment_method"), PaymentMethod)
        if payment_method is None:
            issues.append(
                ValidationIssue("expenses", i, f"payment_method inválido: {values.get('payment_method')!r}.")
            )
            seen_ids.discard(result["id"])
            continue

        if result["type"] == RecurrenceType.FIXED and payment_method != PaymentMethod.CONTA:
            issues.append(ValidationIssue("expenses", i, "Saídas fixas devem ter payment_method=CONTA."))
            seen_ids.discard(result["id"])
            continue

        result["payment_method"] = payment_method
        parsed.append(result)
    return parsed, seen_ids


def _validate_installments(
    rows: list[list], issues: list[ValidationIssue], expenses_by_id: dict[int, dict]
):
    parsed: list[dict] = []
    seen_ids: set[int] = set()
    by_expense: dict[int, list[dict]] = {}

    for i, row in enumerate(rows, start=2):
        values = dict(zip(INSTALLMENTS_COLUMNS, row))

        row_id = _parse_int(values.get("id"))
        if row_id is None or row_id in seen_ids:
            issues.append(ValidationIssue("installments", i, "id inválido, ausente ou duplicado."))
            continue

        expense_id = _parse_int(values.get("expense_id"))
        expense = expenses_by_id.get(expense_id) if expense_id is not None else None
        if expense is None:
            issues.append(ValidationIssue("installments", i, f"expense_id inexistente: {expense_id!r}."))
            continue
        if expense["payment_method"] != PaymentMethod.CREDITO:
            issues.append(
                ValidationIssue(
                    "installments", i, f"expense_id {expense_id} não é uma compra CREDITO."
                )
            )
            continue

        installment_number = _parse_int(values.get("installment_number"))
        total_installments = _parse_int(values.get("total_installments"))
        if (
            installment_number is None
            or total_installments is None
            or not (1 <= total_installments <= 24)
            or not (1 <= installment_number <= total_installments)
        ):
            issues.append(ValidationIssue("installments", i, "installment_number/total_installments inválidos (1-24)."))
            continue

        due_date = _parse_date(values.get("due_date"))
        if due_date is None:
            issues.append(ValidationIssue("installments", i, "due_date inválida."))
            continue

        base_amount = _parse_decimal(values.get("base_amount"))
        if base_amount is None or base_amount <= 0:
            issues.append(ValidationIssue("installments", i, "base_amount inválido ou não positivo."))
            continue

        advanced_to_year = _parse_int(values.get("advanced_to_year"))
        advanced_to_month = _parse_int(values.get("advanced_to_month"))
        if (advanced_to_year is None) != (advanced_to_month is None):
            issues.append(
                ValidationIssue(
                    "installments", i, "advanced_to_year e advanced_to_month devem estar ambos preenchidos ou ambos vazios."
                )
            )
            continue

        seen_ids.add(row_id)
        created_at = _parse_datetime(values.get("created_at")) or datetime.utcnow()
        updated_at = _parse_datetime(values.get("updated_at")) or datetime.utcnow()

        parsed_row = {
            "id": row_id,
            "expense_id": expense_id,
            "installment_number": installment_number,
            "total_installments": total_installments,
            "due_date": due_date,
            "base_amount": base_amount,
            "advanced_to_year": advanced_to_year,
            "advanced_to_month": advanced_to_month,
            "created_at": created_at,
            "updated_at": updated_at,
        }
        parsed.append(parsed_row)
        by_expense.setdefault(expense_id, []).append(parsed_row)

    # Consistency: every CREDITO expense must have installments 1..N with no gaps/dupes,
    # and N must match every sibling's total_installments.
    for expense_id, expense in expenses_by_id.items():
        if expense["payment_method"] != PaymentMethod.CREDITO:
            continue
        siblings = by_expense.get(expense_id, [])
        if not siblings:
            issues.append(
                ValidationIssue("installments", None, f"expense_id {expense_id} (CREDITO) não tem parcelas.")
            )
            continue
        totals = {s["total_installments"] for s in siblings}
        numbers = sorted(s["installment_number"] for s in siblings)
        if len(totals) != 1 or numbers != list(range(1, list(totals)[0] + 1)):
            issues.append(
                ValidationIssue(
                    "installments", None,
                    f"expense_id {expense_id}: parcelas inconsistentes (esperado 1..N sem lacunas/duplicatas).",
                )
            )

    for expense_id, expense in expenses_by_id.items():
        if expense["payment_method"] == PaymentMethod.CONTA and by_expense.get(expense_id):
            issues.append(
                ValidationIssue("installments", None, f"expense_id {expense_id} (CONTA) não deveria ter parcelas.")
            )

    return parsed, seen_ids


def _validate_adjustments(rows: list[list], issues: list[ValidationIssue], valid_installment_ids: set[int]):
    parsed = []
    seen_ids = set()
    for i, row in enumerate(rows, start=2):
        values = dict(zip(ADJUSTMENTS_COLUMNS, row))

        row_id = _parse_int(values.get("id"))
        if row_id is None or row_id in seen_ids:
            issues.append(ValidationIssue("installment_adjustments", i, "id inválido, ausente ou duplicado."))
            continue

        installment_id = _parse_int(values.get("installment_id"))
        if installment_id is None or installment_id not in valid_installment_ids:
            issues.append(
                ValidationIssue(
                    "installment_adjustments", i, f"installment_id inexistente: {installment_id!r}."
                )
            )
            continue

        amount = _parse_decimal(values.get("amount"))
        if amount is None or amount == 0:
            issues.append(ValidationIssue("installment_adjustments", i, "amount inválido ou igual a zero."))
            continue

        description = _parse_str(values.get("description"))
        if description is None:
            issues.append(ValidationIssue("installment_adjustments", i, "description é obrigatória."))
            continue

        seen_ids.add(row_id)
        created_at = _parse_datetime(values.get("created_at")) or datetime.utcnow()
        updated_at = _parse_datetime(values.get("updated_at")) or datetime.utcnow()

        parsed.append(
            {
                "id": row_id,
                "installment_id": installment_id,
                "amount": amount,
                "description": description,
                "created_at": created_at,
                "updated_at": updated_at,
            }
        )
    return parsed, seen_ids


# ---- entry point --------------------------------------------------------------------


def validate_workbook(file_bytes: bytes) -> tuple[ImportPreview, ParsedImport | None]:
    sheets = parse_workbook(file_bytes)
    issues: list[ValidationIssue] = []

    expected_columns = {
        "users": USERS_COLUMNS,
        "entries": ENTRIES_COLUMNS,
        "expenses": EXPENSES_COLUMNS,
        "installments": INSTALLMENTS_COLUMNS,
        "installment_adjustments": ADJUSTMENTS_COLUMNS,
    }
    headers_ok = True
    for name, columns in expected_columns.items():
        header, _data = sheets[name]
        if not _check_header(name, header, columns, issues):
            headers_ok = False

    if not headers_ok:
        preview = ImportPreview(valid=False, errors=issues, warnings=[], counts={})
        return preview, None

    users, valid_user_ids = _validate_users(sheets["users"][1], issues)
    entries, _ = _validate_entries(sheets["entries"][1], issues, valid_user_ids)
    expenses, _ = _validate_expenses(sheets["expenses"][1], issues, valid_user_ids)
    expenses_by_id = {e["id"]: e for e in expenses}
    installments, valid_installment_ids = _validate_installments(
        sheets["installments"][1], issues, expenses_by_id
    )
    adjustments, _ = _validate_adjustments(
        sheets["installment_adjustments"][1], issues, valid_installment_ids
    )

    counts = {
        "users": len(users),
        "entries": len(entries),
        "expenses": len(expenses),
        "installments": len(installments),
        "installment_adjustments": len(adjustments),
    }

    valid = len(issues) == 0
    preview = ImportPreview(valid=valid, errors=issues, warnings=[], counts=counts)
    parsed = (
        ParsedImport(
            users=users,
            entries=entries,
            expenses=expenses,
            installments=installments,
            installment_adjustments=adjustments,
        )
        if valid
        else None
    )
    return preview, parsed


@dataclass
class ImportResult:
    counts: dict[str, int]
    temporary_passwords: dict[str, str]  # email -> plaintext password, returned ONCE


def execute_import(engine: Engine, parsed: ParsedImport) -> ImportResult:
    temporary_passwords: dict[str, str] = {}

    with engine.begin() as conn:
        conn.execute(text("SET FOREIGN_KEY_CHECKS=0"))

        # refresh_tokens has no ORM model of its own (see app/models/refresh_token.py);
        # it's deleted here alongside users since every row references a user_id that's
        # about to be wiped. Same for the Assistente's chat history: with FK checks off,
        # ON DELETE CASCADE doesn't fire, and a conversation (or a saved LLM API key)
        # left behind would belong to whichever imported user ends up with the same id.
        for table_name in (
            "assistant_settings",
            "chat_messages",
            "chat_conversations",
            InstallmentAdjustment.__table__.name,
            Installment.__table__.name,
            Expense.__table__.name,
            Entry.__table__.name,
            "refresh_tokens",
            User.__table__.name,
        ):
            conn.execute(text(f"DELETE FROM {table_name}"))

        for row in parsed.users:
            temp_password = generate_temp_password()
            temporary_passwords[row["email"]] = temp_password
            conn.execute(
                User.__table__.insert().values(
                    id=row["id"],
                    name=row["name"],
                    email=row["email"],
                    password_hash=hash_password(temp_password),
                    photo=row["photo"],
                    role=row["role"],
                    created_at=row["created_at"],
                    updated_at=row["updated_at"],
                )
            )

        if parsed.entries:
            conn.execute(Entry.__table__.insert(), parsed.entries)
        if parsed.expenses:
            conn.execute(Expense.__table__.insert(), parsed.expenses)
        if parsed.installments:
            conn.execute(Installment.__table__.insert(), parsed.installments)
        if parsed.installment_adjustments:
            conn.execute(InstallmentAdjustment.__table__.insert(), parsed.installment_adjustments)

        # No manual AUTO_INCREMENT reset needed: InnoDB automatically advances a table's
        # counter past any explicitly-inserted id that meets or exceeds it. (Deliberately
        # NOT using ALTER TABLE here even as a belt-and-suspenders measure - it's DDL,
        # and DDL causes an implicit commit in MySQL, which would silently break the
        # all-or-nothing rollback guarantee for everything inserted above.)

        conn.execute(text("SET FOREIGN_KEY_CHECKS=1"))

    return ImportResult(
        counts={
            "users": len(parsed.users),
            "entries": len(parsed.entries),
            "expenses": len(parsed.expenses),
            "installments": len(parsed.installments),
            "installment_adjustments": len(parsed.installment_adjustments),
        },
        temporary_passwords=temporary_passwords,
    )
